# -*- coding: utf-8 -*-
import json
from pathlib import Path

try:
    import openpyxl
except Exception as e:
    print("openpyxl error", e)
    openpyxl = None

try:
    import xlrd
except Exception as e:
    print("xlrd error", e)
    xlrd = None

out_dir = Path(r"C:\Users\Renzo\Desktop\WEB MEMORIAS DESCRIPTIVAS")
files = [
    r"D:\Renzo\RENZO\6. EXCELS\12. HIDRAULICA\2. DRENAJE PLUVIAL\DISEÑO CANALETA.xlsx",
    r"D:\Renzo\RENZO\6. EXCELS\12. HIDRAULICA\SIFON INVERTIDO.xls",
    r"D:\Renzo\RENZO\6. EXCELS\12. HIDRAULICA\SIFON.xls",
    r"D:\Renzo\RENZO\6. EXCELS\12. HIDRAULICA\f314628096_sifon.xls",
    r"D:\Renzo\RENZO\6. EXCELS\12. HIDRAULICA\DISEÑO DE CUNETAS.xls",
    r"D:\Renzo\RENZO\6. EXCELS\12. HIDRAULICA\ALCANTARILLA CAJON.xlsx",
    r"D:\Renzo\RENZO\6. EXCELS\12. HIDRAULICA\EMBALSES.xls",
    r"D:\Renzo\RENZO\6. EXCELS\12. HIDRAULICA\f314591232_ALIVIADERO.xls",
    r"D:\Renzo\RENZO\6. EXCELS\12. HIDRAULICA\f314646528_desarenador.xls",
    r"D:\Renzo\RENZO\6. EXCELS\12. HIDRAULICA\GALERIA FILTRANTE.xls",
    r"D:\Renzo\RENZO\6. EXCELS\12. HIDRAULICA\CRUCE AEREO.xlsx",
    r"D:\Renzo\RENZO\6. EXCELS\12. HIDRAULICA\MC TOMA.xlsx",
]


def cell_str(v):
    if v is None:
        return ""
    if isinstance(v, float):
        if v == int(v):
            return str(int(v))
        return f"{v:.8g}"
    return str(v).strip()


def dump_xlsx(path, max_rows=200, max_cols=18):
    result = {"file": str(path), "sheets": []}
    wb_val = openpyxl.load_workbook(path, data_only=True, read_only=True)
    wb_f = openpyxl.load_workbook(path, data_only=False, read_only=True)
    n_images = 0
    try:
        wb_img = openpyxl.load_workbook(path)
        for name in wb_img.sheetnames:
            ws = wb_img[name]
            imgs = getattr(ws, "_images", []) or []
            n_images += len(imgs)
    except Exception:
        pass

    for name in wb_val.sheetnames:
        ws_v = wb_val[name]
        ws_f = wb_f[name]
        rows = []
        formulas = []
        for i, (row_v, row_f) in enumerate(
            zip(
                ws_v.iter_rows(max_row=max_rows, max_col=max_cols, values_only=True),
                ws_f.iter_rows(max_row=max_rows, max_col=max_cols, values_only=True),
            )
        ):
            vals = [cell_str(c) for c in row_v]
            fmls = [cell_str(c) for c in row_f]
            if any(vals) or any(fmls):
                rows.append({"r": i + 1, "v": vals})
                formulas.append({"r": i + 1, "f": fmls})
        result["sheets"].append(
            {"name": name, "rows": rows, "formulas": formulas, "images": n_images}
        )
    result["images_total"] = n_images
    return result


def dump_xls(path, max_rows=200, max_cols=18):
    result = {"file": str(path), "sheets": []}
    book = xlrd.open_workbook(path, formatting_info=False)
    for name in book.sheet_names():
        sh = book.sheet_by_name(name)
        rows = []
        formulas = []
        nr = min(sh.nrows, max_rows)
        nc = min(sh.ncols, max_cols)
        for i in range(nr):
            vals = []
            fmls = []
            for j in range(nc):
                cell = sh.cell(i, j)
                vals.append(cell_str(cell.value))
                fmls.append(cell_str(cell.value))
            if any(vals):
                rows.append({"r": i + 1, "v": vals})
                formulas.append({"r": i + 1, "f": fmls})
        result["sheets"].append({"name": name, "rows": rows, "formulas": formulas, "images": 0})
    return result


all_out = []
for f in files:
    p = Path(f)
    print("===", p.name, "exists=", p.exists(), "size=", p.stat().st_size if p.exists() else 0)
    if not p.exists():
        all_out.append({"file": str(p), "error": "missing"})
        continue
    try:
        if p.suffix.lower() == ".xlsx":
            all_out.append(dump_xlsx(p))
        else:
            all_out.append(dump_xls(p))
        sheets = all_out[-1].get("sheets", [])
        print("  sheets:", [s["name"] for s in sheets], "nrows0=", len(sheets[0]["rows"]) if sheets else 0)
    except Exception as e:
        print("  ERROR", type(e).__name__, e)
        all_out.append({"file": str(p), "error": f"{type(e).__name__}: {e}"})

out_path = out_dir / "_hidro_extract.json"
out_path.write_text(json.dumps(all_out, ensure_ascii=False, indent=1), encoding="utf-8")
print("WROTE", out_path, "bytes", out_path.stat().st_size)
