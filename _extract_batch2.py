import json
from pathlib import Path

try:
    import openpyxl
except Exception as e:
    print("openpyxl", e)
try:
    import xlrd
except Exception as e:
    print("xlrd", e)

out_dir = Path(r"C:\Users\Renzo\Desktop\WEB MEMORIAS DESCRIPTIVAS")
files = [
    r"d:\Renzo\RENZO\6. EXCELS\1. ESTRUCTURAS\1. EDIFIC\7. PREDIMENSIONAMIENTO\PRED COL VIG.xlsx",
    r"d:\Renzo\RENZO\6. EXCELS\1. ESTRUCTURAS\1. EDIFIC\7. PREDIMENSIONAMIENTO\PRED ELEM ESTR.xls",
    r"d:\Renzo\RENZO\6. EXCELS\1. ESTRUCTURAS\1. EDIFIC\7. PREDIMENSIONAMIENTO\PREDIM COL.xlsx",
    r"d:\Renzo\RENZO\6. EXCELS\1. ESTRUCTURAS\1. EDIFIC\8. LOSAS\LOS ARMADA.xlsx",
    r"d:\Renzo\RENZO\6. EXCELS\1. ESTRUCTURAS\1. EDIFIC\8. LOSAS\LOSA ARMADA 1DIR.xls",
    r"d:\Renzo\RENZO\6. EXCELS\1. ESTRUCTURAS\1. EDIFIC\8. LOSAS\Diseño de Viguetas.xlsx",
    r"d:\Renzo\RENZO\6. EXCELS\1. ESTRUCTURAS\1. EDIFIC\8. LOSAS\DIS LOSA MZ.xls",
    r"d:\Renzo\RENZO\6. EXCELS\1. ESTRUCTURAS\1. EDIFIC\8. LOSAS\ANALISIS TECHO T CANOPY.xlsx",
    r"d:\Renzo\RENZO\6. EXCELS\1. ESTRUCTURAS\1. EDIFIC\9. LOSA NERVADA\f314573824.xls",
    r"d:\Renzo\RENZO\6. EXCELS\1. ESTRUCTURAS\1. EDIFIC\10. DISE MUROS C°\DISEÑO DE PLACAS.xlsx",
    r"d:\Renzo\RENZO\6. EXCELS\5. INST ELECTRICAS\CAL INST ELECTRIC.xlsx",
    r"d:\Renzo\RENZO\6. EXCELS\5. INST ELECTRICAS\Cálculo de Máxima Demanda y Potencia Instalada (1).xlsx",
    r"d:\Renzo\RENZO\6. EXCELS\6. TASACIONES\Tasaciones.xlsx",
    r"d:\Renzo\RENZO\6. EXCELS\6. TASACIONES\TASACION DE MAQ.xlsx",
    r"d:\Renzo\RENZO\6. EXCELS\7. SUELOS\CAP PORTANTE.xlsx",
    r"d:\Renzo\RENZO\6. EXCELS\7. SUELOS\CALCULO DE PRESION DE TERRENO.xlsx",
    r"d:\Renzo\RENZO\6. EXCELS\7. SUELOS\C HUMEDAD.xlsx",
    r"d:\Renzo\RENZO\6. EXCELS\7. SUELOS\CBR (C2-M1) okk.xlsx",
    r"d:\Renzo\RENZO\6. EXCELS\7. SUELOS\ESTUDIO DE SUELOS COMPLETO (2).xlsm",
    r"d:\Renzo\RENZO\6. EXCELS\8. AGUA Y ALCANTARILLADO\ALCANTARILLA TIPO CAJON.xlsx",
    r"d:\Renzo\RENZO\6. EXCELS\8. AGUA Y ALCANTARILLADO\TANQUE SEPTICO.xls",
    r"d:\Renzo\RENZO\6. EXCELS\9. PAVIMENTOS\CAMINO VECINAL.xlsx",
    r"d:\Renzo\RENZO\6. EXCELS\9. PAVIMENTOS\PAVIMENTO PARA AVIACION.xls",
    r"d:\Renzo\RENZO\6. EXCELS\10. EXCLS UNI\DISEÑO DE VIGAS UNI.xls",
    r"d:\Renzo\RENZO\6. EXCELS\10. EXCLS UNI\Diagrama M Viga.xls",
    r"d:\Renzo\RENZO\6. EXCELS\10. EXCLS UNI\MET CROSS.xls",
    r"d:\Renzo\RENZO\6. EXCELS\10. EXCLS UNI\f315326464_BeamAnal.xls",
]

def cell_str(v):
    if v is None:
        return ""
    if isinstance(v, float):
        if v == int(v):
            return str(int(v))
        return f"{v:.6g}"
    return str(v).strip()

def dump_xlsx(path, max_rows=70, max_cols=14):
    wb = openpyxl.load_workbook(path, data_only=True, read_only=True)
    sheets = []
    for name in wb.sheetnames[:8]:
        ws = wb[name]
        rows = []
        for i, row in enumerate(ws.iter_rows(max_row=max_rows, max_col=max_cols, values_only=True)):
            vals = [cell_str(c) for c in row]
            if any(vals):
                rows.append(vals)
        sheets.append({"name": name, "rows": rows[:50], "images": 0})
    wb.close()
    try:
        wb2 = openpyxl.load_workbook(path, data_only=True)
        for name in wb2.sheetnames[:8]:
            ws = wb2[name]
            n = len(getattr(ws, "_images", []) or [])
            for s in sheets:
                if s["name"] == name:
                    s["images"] = n
        wb2.close()
    except Exception:
        pass
    return sheets, list(openpyxl.load_workbook(path, read_only=True).sheetnames)

def dump_xls(path, max_rows=70, max_cols=14):
    book = xlrd.open_workbook(path)
    sheets = []
    for sh in book.sheets()[:8]:
        rows = []
        rmax = min(sh.nrows, max_rows)
        cmax = min(sh.ncols, max_cols)
        for r in range(rmax):
            vals = [cell_str(sh.cell_value(r, c)) for c in range(cmax)]
            if any(vals):
                rows.append(vals)
        sheets.append({"name": sh.name, "rows": rows[:50], "images": 0})
    return sheets, [s.name for s in book.sheets()]

report = []
for f in files:
    p = Path(f)
    item = {"file": p.name, "path": str(p), "exists": p.exists(), "size": p.stat().st_size if p.exists() else 0}
    print("FILE", p.name, "exists", p.exists())
    if not p.exists():
        item["error"] = "NOT FOUND"
        report.append(item)
        continue
    try:
        ext = p.suffix.lower()
        if ext in (".xlsx", ".xlsm"):
            sheets, names = dump_xlsx(p)
        else:
            sheets, names = dump_xls(p)
        item["sheet_names"] = names
        item["sheets"] = sheets
        item["n_images"] = sum(s["images"] for s in sheets)
    except Exception as e:
        item["error"] = f"{type(e).__name__}: {e}"
        print(" ERR", item["error"])
    report.append(item)

out = out_dir / "_excel_batch2.json"
with open(out, "w", encoding="utf-8") as fh:
    json.dump(report, fh, ensure_ascii=False, indent=1)

txt = out_dir / "_excel_batch2.txt"
lines = []
for item in report:
    lines.append("=" * 80)
    lines.append(f"{item['file']} exists={item['exists']} imgs={item.get('n_images')} err={item.get('error','')}")
    lines.append(f"SHEETS: {item.get('sheet_names')}")
    for s in item.get("sheets", [])[:5]:
        lines.append(f"  -- {s['name']} imgs {s['images']}")
        n = 0
        for row in s["rows"]:
            text = " | ".join([c for c in row if c])
            if len(text) < 4:
                continue
            lines.append("     " + text[:220])
            n += 1
            if n >= 28:
                break
with open(txt, "w", encoding="utf-8") as fh:
    fh.write("\n".join(lines))
print("WROTE", txt)
