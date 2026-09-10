# -*- coding: utf-8 -*-
import os
from openpyxl import load_workbook

folder = r"d:\Renzo\RENZO\6. EXCELS\6. TASACIONES"
out = r"C:\Users\Renzo\Desktop\WEB MEMORIAS DESCRIPTIVAS\_tas_extract.txt"

files = [
    "TASACION DE TERRENO ACTUALIZADA.xlsx",
    "Tasaciones.xlsx",
    "Tasaciones2.xlsx",
    "TASACION DE MAQ.xlsx",
    "f33522688.xlsx",
]

def dump_xlsx(path, fh, max_rows=80, max_cols=14):
    fh.write(f"\n\n========== {os.path.basename(path)} ==========\n")
    wb = load_workbook(path, data_only=True, read_only=True)
    fh.write("SHEETS: " + " | ".join(wb.sheetnames) + "\n")
    for name in wb.sheetnames:
        ws = wb[name]
        fh.write(f"\n----- hoja: {name} -----\n")
        n = 0
        for row in ws.iter_rows(max_row=max_rows, max_col=max_cols, values_only=True):
            vals = ["" if c is None else str(c).replace("\n", " ").strip() for c in row]
            if any(vals):
                fh.write(" | ".join(vals) + "\n")
                n += 1
            if n > 70:
                fh.write("... (truncated)\n")
                break
    wb.close()

with open(out, "w", encoding="utf-8") as fh:
    for f in files:
        p = os.path.join(folder, f)
        if os.path.exists(p):
            try:
                dump_xlsx(p, fh)
            except Exception as e:
                fh.write(f"ERROR {f}: {e}\n")
        else:
            fh.write(f"MISSING {f}\n")

print("wrote", out)
