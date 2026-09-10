# -*- coding: utf-8 -*-
from openpyxl import load_workbook
from pypdf import PdfReader
import os

folder = r"d:\Renzo\RENZO\6. EXCELS\6. TASACIONES"
out = r"C:\Users\Renzo\Desktop\WEB MEMORIAS DESCRIPTIVAS\_tas_extract2.txt"

def dump_sheet(path, sheet, max_row=200, max_col=18):
    wb = load_workbook(path, data_only=True, read_only=True)
    ws = wb[sheet]
    lines = [f"===== {os.path.basename(path)} :: {sheet} ====="]
    n = 0
    for row in ws.iter_rows(max_row=max_row, max_col=max_col, values_only=True):
        vals = ["" if c is None else str(c).replace("\n", " ").strip() for c in row]
        if any(vals):
            lines.append(" | ".join(vals))
            n += 1
    wb.close()
    return "\n".join(lines)

chunks = []
p = os.path.join(folder, "TASACION DE TERRENO ACTUALIZADA.xlsx")
chunks.append(dump_sheet(p, "EDIFICACIONES", 180, 16))
chunks.append(dump_sheet(p, "TERRENO", 180, 16))
chunks.append(dump_sheet(p, "Inicio", 80, 16))

# formulas sheet names already known
p2 = os.path.join(folder, "TASACION DE MAQ.xlsx")
wb = load_workbook(p2, data_only=True, read_only=True)
chunks.append("MAQ SHEETS: " + " | ".join(wb.sheetnames))
wb.close()
chunks.append(dump_sheet(p2, load_workbook(p2, read_only=True).sheetnames[0], 60, 12))

pdf_name = None
for f in os.listdir(folder):
    if "PAUL" in f.upper() or "HARRIS" in f.upper():
        pdf_name = f
        break
    if "TASACI" in f.upper() and f.lower().endswith(".pdf") and "VIVIENDA" in f.upper():
        pdf_name = f

chunks.append("\nPDF found: " + str(pdf_name))
if pdf_name:
    reader = PdfReader(os.path.join(folder, pdf_name))
    chunks.append(f"pages={len(reader.pages)}")
    for i, page in enumerate(reader.pages[:12]):
        t = page.extract_text() or ""
        chunks.append(f"\n----- PDF p{i+1} -----\n{t[:4500]}")

with open(out, "w", encoding="utf-8") as fh:
    fh.write("\n\n".join(chunks))
print("wrote", out, "chars", sum(len(c) for c in chunks))
