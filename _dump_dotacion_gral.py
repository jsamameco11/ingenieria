# -*- coding: utf-8 -*-
import json
import sys
from pathlib import Path

sys.stdout.reconfigure(encoding="utf-8")
from openpyxl import load_workbook

fp = Path(r"d:\Renzo\RENZO\6. EXCELS\15. INST SANITARIAS\CALCULO DE DOTACION DE AGUA.xlsx")
print("exists", fp.exists(), "size", fp.stat().st_size)
wb = load_workbook(fp, data_only=True, read_only=True)
print("sheets:", wb.sheetnames)
for name in wb.sheetnames:
    ws = wb[name]
    print("\n" + "=" * 80)
    print("SHEET", name)
    n = 0
    for i, row in enumerate(ws.iter_rows(max_row=120, max_col=12, values_only=True), 1):
        vals = [str(c).strip() if c is not None else "" for c in row]
        if any(vals):
            print(f"r{i:3d}: " + " | ".join(x for x in vals if x)[:200])
            n += 1
    print(f"[{n} nonempty]")
