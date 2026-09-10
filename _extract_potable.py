# -*- coding: utf-8 -*-
import json
import sys
from pathlib import Path

sys.stdout.reconfigure(encoding="utf-8")

import openpyxl
import xlrd

out_dir = Path(r"C:\Users\Renzo\Desktop\WEB MEMORIAS DESCRIPTIVAS")
base = Path(r"d:\Renzo\RENZO\6. EXCELS\17. SIST TRATAM A POTABLE")
files = [
    "DOTACION DE AGUA .xlsx",
    "f314812416.xls",
    "f314852352.xls",
    "f314853376.xls",
    "f314854400.xls",
    "f314855424.xls",
    "f314856448.xls",
    "f314857472.xls",
    "f314858496.xls",
    "f314859520.xls",
    "f315129856.xls",
    "f315431936.xls",
    "f315509760_WEIGHTS_Program.xls",
    "L IMPULSION.xlsx",
    "SISTEMA DE CLORACION.xlsx",
]


def cell_str(v):
    if v is None:
        return ""
    if isinstance(v, float):
        if v == int(v):
            return str(int(v))
        return f"{v:.8g}"
    return str(v).strip()


def dump_xlsx(path, max_rows=250, max_cols=16):
    result = {"file": path.name, "kind": "xlsx", "sheets": []}
    wb_v = openpyxl.load_workbook(path, data_only=True, read_only=True)
    wb_f = openpyxl.load_workbook(path, data_only=False, read_only=True)
    n_images = 0
    try:
        wb_img = openpyxl.load_workbook(path)
        for name in wb_img.sheetnames:
            n_images += len(getattr(wb_img[name], "_images", []) or [])
    except Exception:
        pass
    for name in wb_v.sheetnames:
        ws_v = wb_v[name]
        ws_f = wb_f[name]
        rows, formulas = [], []
        for i, (row_v, row_f) in enumerate(
            zip(
                ws_v.iter_rows(max_row=max_rows, max_col=max_cols, values_only=True),
                ws_f.iter_rows(max_row=max_rows, max_col=max_cols, values_only=True),
            )
        ):
            vals = [cell_str(c) for c in row_v]
            fmls = [cell_str(c) for c in row_f]
            if any(vals) or any(fmls):
                rows.append({"r": i + 1, "v": vals})
                formulas.append({"r": i + 1, "f": fmls})
        result["sheets"].append({"name": name, "rows": rows, "formulas": formulas})
    result["images"] = n_images
    return result


def dump_xls(path, max_rows=250, max_cols=16):
    result = {"file": path.name, "kind": "xls", "sheets": []}
    book = xlrd.open_workbook(str(path), formatting_info=False)
    for name in book.sheet_names():
        sh = book.sheet_by_name(name)
        rows = []
        nr = min(sh.nrows, max_rows)
        nc = min(sh.ncols, max_cols)
        for i in range(nr):
            vals = [cell_str(sh.cell_value(i, j)) for j in range(nc)]
            if any(vals):
                rows.append({"r": i + 1, "v": vals})
        result["sheets"].append({"name": name, "nrows": sh.nrows, "ncols": sh.ncols, "rows": rows})
    result["images"] = 0
    return result


all_out = []
for name in files:
    p = base / name
    print("=" * 80)
    print(p.name, "exists=", p.exists(), "size=", p.stat().st_size if p.exists() else 0)
    if not p.exists():
        all_out.append({"file": name, "error": "missing"})
        continue
    try:
        data = dump_xlsx(p) if p.suffix.lower() == ".xlsx" else dump_xls(p)
        all_out.append(data)
        print("  sheets:", [s["name"] for s in data["sheets"]])
        for s in data["sheets"]:
            print(f"  -- {s['name']} nonempty={len(s['rows'])}")
            for row in s["rows"][:12]:
                joined = " | ".join(x for x in row["v"] if x)[:180]
                if joined:
                    print(f"     r{row['r']}: {joined}")
    except Exception as e:
        print("  ERROR", type(e).__name__, e)
        all_out.append({"file": name, "error": f"{type(e).__name__}: {e}"})

out_path = out_dir / "_potable_extract.json"
out_path.write_text(json.dumps(all_out, ensure_ascii=False, indent=1), encoding="utf-8")
print("WROTE", out_path, "bytes", out_path.stat().st_size)
