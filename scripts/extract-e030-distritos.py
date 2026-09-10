import json
import os
import re
from collections import Counter, defaultdict

import openpyxl

src = r"d:\Renzo\RENZO\6. EXCELS\1. ESTRUCTURAS\1. EDIFIC\4. EDIFICACIONES COMPLETO\ANA DISEÑO EDIFI.xlsx"
out = r"C:\Users\Renzo\Desktop\WEB MEMORIAS DESCRIPTIVAS\src\lib\e030\distritos.json"

wb = openpyxl.load_workbook(src, data_only=True)
ws = wb["Cuadros E030"]

SMALL = {"de", "del", "la", "las", "los", "y", "e", "da", "do"}


def parse_ref(ref: str):
    ref = ref.replace("'", "")
    m = re.search(r"\$([A-Z]+)\$(\d+):\$([A-Z]+)\$(\d+)", ref)
    if m:
        return m.group(1), int(m.group(2)), m.group(3), int(m.group(4))
    m = re.search(r"\$([A-Z]+)\$(\d+)", ref)
    if m:
        return m.group(1), int(m.group(2)), m.group(1), int(m.group(2))
    return None


def col_idx(col: str) -> int:
    n = 0
    for ch in col:
        n = n * 26 + ord(ch) - 64
    return n


def nice(code: str) -> str:
    s = re.sub(r"^O\d+_", "", code or "")
    s = s.rstrip(".").replace("_", " ").replace(".", " ")
    s = re.sub(r"\s+", " ", s).strip()
    parts = s.split(" ")
    out = []
    for i, p in enumerate(parts):
        w = p.lower()
        if i > 0 and w in SMALL:
            out.append(w)
        elif not w:
            continue
        else:
            out.append(w[0].upper() + w[1:])
    return " ".join(out)


def dist_name(raw: str) -> str:
    s = re.sub(r"^\d+\s+", "", (raw or "").strip())
    s = re.sub(r"\s+", " ", s)
    parts = s.split(" ")
    out = []
    for i, p in enumerate(parts):
        w = p.lower()
        if i > 0 and w in SMALL:
            out.append(w)
        elif not w:
            continue
        else:
            out.append(w[0].upper() + w[1:])
    return " ".join(out)


def mode_zone(zs: list[int]) -> int | None:
    valid = [z for z in zs if z in (1, 2, 3, 4)]
    if not valid:
        return None
    counts = Counter(valid)
    return counts.most_common(1)[0][0]


depts: dict = {}
provs: dict = {}
for name, dn in wb.defined_names.items():
    attr = str(dn.attr_text or "")
    if "Cuadros E030" not in attr:
        continue
    parsed = parse_ref(attr)
    if not parsed:
        continue
    c1, r1, c2, r2 = parsed
    ci = col_idx(c1)
    vals = []
    for r in range(r1, r2 + 1):
        v = ws.cell(r, ci).value
        if v is not None and str(v).strip():
            vals.append((r, str(v).strip()))
    if c1 == "B":
        depts[name] = {"name": nice(name), "provinces": [v for _, v in vals]}
    elif c1 == "C":
        rows = []
        for r, v in vals:
            zraw = ws.cell(r, 4).value
            try:
                z = int(zraw)
            except (TypeError, ValueError):
                z = None
            ambito = str(ws.cell(r, 5).value or "").strip() or "TODOS LOS DISTRITOS"
            rows.append({"name": dist_name(v), "z": z, "ambito": ambito})
        provs[name] = {"name": nice(name), "rows": rows}

# Uniformar zona dentro de cada bloque de ámbito (corrige arrastre 5, 6, 7…)
for pr in provs.values():
    groups: dict[str, list] = defaultdict(list)
    for row in pr["rows"]:
        groups[row["ambito"]].append(row)
    for block in groups.values():
        z = mode_zone([r["z"] for r in block if r["z"] is not None])
        if z is None:
            continue
        for r in block:
            r["z"] = z

tree = []
for _, d in sorted(depts.items(), key=lambda x: x[1]["name"]):
    plist = []
    for pcode in d["provinces"]:
        pr = provs.get(pcode)
        if not pr:
            continue
        dists = [{"n": x["name"], "z": x["z"]} for x in pr["rows"] if x["z"] in (1, 2, 3, 4) and x["name"]]
        # quitar duplicados conservando orden
        seen = set()
        uniq = []
        for x in dists:
            k = x["n"].casefold()
            if k in seen:
                continue
            seen.add(k)
            uniq.append(x)
        if uniq:
            plist.append({"n": pr["name"], "d": uniq})
    if plist:
        tree.append({"n": d["name"], "p": plist})

os.makedirs(os.path.dirname(out), exist_ok=True)
with open(out, "w", encoding="utf-8") as f:
    json.dump(tree, f, ensure_ascii=False, separators=(",", ":"))

zonas = Counter(x["z"] for d in tree for p in d["p"] for x in p["d"])
print("depts", len(tree))
print("provs", sum(len(d["p"]) for d in tree))
print("dist", sum(len(p["d"]) for d in tree for p in d["p"]))
print("zonas", dict(sorted(zonas.items())))
print("bytes", os.path.getsize(out))
lima = next(d for d in tree if d["n"] == "Lima")
print("Lima provs", [p["n"] for p in lima["p"]])
lp = next(p for p in lima["p"] if p["n"] == "Lima")
print("Lima dist sample", lp["d"][:8], "... n=", len(lp["d"]))
print("Miraflores", [x for x in lp["d"] if "Miraflores" in x["n"]])
