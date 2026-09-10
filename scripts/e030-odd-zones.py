import re
import sys
import openpyxl

sys.stdout.reconfigure(encoding="utf-8")
src = r"d:\Renzo\RENZO\6. EXCELS\1. ESTRUCTURAS\1. EDIFIC\4. EDIFICACIONES COMPLETO\ANA DISEÑO EDIFI.xlsx"
wb = openpyxl.load_workbook(src, data_only=True)
ws = wb["Cuadros E030"]
print("odd zones:")
for r in range(33, ws.max_row + 1):
    v = ws.cell(r, 3).value
    d = ws.cell(r, 4).value
    if isinstance(v, str) and re.match(r"^\d+\s+", v.strip()):
        try:
            z = int(d)
        except (TypeError, ValueError):
            z = None
        if z is not None and z not in (1, 2, 3, 4):
            print(r, v, d, ws.cell(r, 5).value, ws.cell(r, 2).value)
