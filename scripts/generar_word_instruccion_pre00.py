# -*- coding: utf-8 -*-
"""Una sola instrucción PRE-00: protocolo + catálogo completo por especialidad y por plantilla."""
from __future__ import annotations

import re
from collections import defaultdict
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

ROOT = Path(r"C:\Users\Renzo\Desktop\WEB MEMORIAS DESCRIPTIVAS")
LIB = ROOT / "src" / "lib" / "presupuesto"

OUTS = [
    ROOT / "docs" / "INSTRUCCION_AGENTE_PRE00_METRADO.docx",
    Path(
        r"D:\Renzo\RENZO\7. ALCOSAVI\COTIZACIONES Y CONTRATOS 2026"
        r"\01. CISSAC - Conjunto Residencial Los Ciruelos II"
        r"\02_N68_IIEE_Instalaciones_electricas"
        r"\02_LISTA_Y_CRONOGRAMA_MATERIALES"
        r"\15_INSTRUCCION_AGENTE_PRE00_PDF_POR_ESPECIALIDAD.docx"
    ),
]

NAVY = RGBColor(0x12, 0x26, 0x3A)
GOLD = RGBColor(0x8A, 0x6A, 0x32)
RED = RGBColor(0x8A, 0x2C, 0x2C)
INK = RGBColor(0x1F, 0x1A, 0x14)

STR = r'"(?:[^"\\]|\\.)*"'
P_RE = re.compile(
    rf"p\(\s*({STR})\s*,\s*({STR})\s*,\s*({STR})\s*,\s*({STR})\s*,\s*({STR})"
)


def unquote(s: str) -> str:
    return s[1:-1].replace('\\"', '"').replace("\\'", "'").replace("\\\\", "\\")
THW_RE = re.compile(r'thw\("(\d+)",\s*"([^"]+)"')
AE_RE = re.compile(r'\{\s*n:\s*"(\d+)",\s*nombre:\s*"([^"]+)"')
AE_DESDE_RE = re.compile(r'"(\d{2})":\s*"(I-\d|II-[12E]|III-[12E])"')
COD_DESDE_RE = re.compile(r'"(I[EM]-0[6-9]\.[0-9.]+)":\s*"(I-\d|II-[12E]|III-[12E])"')
EQ_RE = re.compile(
    r'e\("((?:EQ|MA|UT|MO|HE)-[A-Z]+-\d{3})",\s*"([^"]+)",\s*[\d.]+,\s*[\d.]+,\s*"(\w+)",\s*"([^"]+)",\s*"([^"]+)",\s*"([^"]+)",\s*"(I-1|I-2|I-3|I-4|II-1|II-2|II-E|III-1|III-E|III-2)"'
)
META_RE = re.compile(
    r'"(I-\d|II-[12E]|III-[12E])":\s*\{.*?nombre:\s*"([^"]+)".*?alias:\s*"([^"]+)".*?resumen:\s*"([^"]+)".*?norma:\s*"([^"]+)"',
    re.S,
)

RANK = {
    "I-1": 1,
    "I-2": 2,
    "I-3": 3,
    "I-4": 4,
    "II-1": 5,
    "II-2": 6,
    "II-E": 6,
    "III-1": 7,
    "III-E": 7,
    "III-2": 8,
}

CATS = list(RANK)

ESP_LABEL = {
    "arquitectura": "Arquitectura (ARQ-)",
    "estructuras": "Estructuras (EST-)",
    "sanitarias": "Instalaciones sanitarias (IS-)",
    "electricas": "Instalaciones eléctricas (IE-)",
    "comunicaciones": "Comunicaciones (COM-)",
    "mecanicas": "Instalaciones mecánicas (IM-)",
    "equipamiento": "Equipamiento hospitalario (EQ-/MA-/UT-/MO-/HE-)",
    "electromecanicas": "Instalaciones electromecánicas (IEM-)",
    "pavimentos": "Pavimentación (P-)",
    "saneamiento": "Saneamiento (S-)",
    "carreteras": "Carreteras y puentes (CAR-)",
    "hidraulica": "Hidráulica (HID-)",
    "habilitaciones": "Habilitaciones urbanas (HAB-)",
}

# Prefijos que la plantilla PUEDE usar. El plano manda; esto no es semilla de metrado.
PLANTILLA_CATALOGO = [
    (
        "Edificaciones",
        [
            ("Vivienda unifamiliar", "unifamiliar", ["ARQ-", "EST-", "IS-01", "IS-02", "IS-03", "IS-04", "IS-05", "IE-01", "IE-02", "IE-03", "IE-04", "IE-05", "COM-", "IM-01", "IM-05"]),
            ("Edificio multifamiliar", "multifamiliar", ["ARQ-", "EST-", "IS-", "IE-", "COM-", "IM-01", "IM-02", "IM-03", "IM-05", "IM-06", "IEM-"]),
            ("Hotel / hospedaje", "hotel", ["ARQ-", "EST-", "IS-", "IE-", "COM-", "IM-01", "IM-02", "IM-03", "IM-05", "IEM-"]),
            ("Restaurante", "restaurante", ["ARQ-", "EST-", "IS-", "IE-", "COM-", "IM-01", "IM-02", "IM-05", "IM-06"]),
            ("Colegio", "colegio", ["ARQ-", "EST-", "IS-", "IE-", "COM-", "IM-01", "IM-02", "IM-05"]),
            ("Oficinas", "oficinas", ["ARQ-", "EST-", "IS-", "IE-", "COM-", "IM-01", "IM-02", "IM-03", "IM-05", "IEM-"]),
            ("Centro comercial", "centro-comercial", ["ARQ-", "EST-", "IS-", "IE-", "COM-", "IM-01", "IM-02", "IM-03", "IM-05", "IEM-"]),
            ("Mercado", "mercado", ["ARQ-", "EST-", "IS-", "IE-", "COM-", "IM-01", "IM-02", "IM-05"]),
            ("Local comunal", "local-comunal", ["ARQ-", "EST-", "IS-", "IE-", "COM-", "IM-01", "IM-02", "IM-05"]),
            ("Comisaría", "comisaria", ["ARQ-", "EST-", "IS-", "IE-", "COM-", "IM-01", "IM-02", "IM-03", "IM-05"]),
        ],
    ),
    (
        "Puentes",
        [
            ("Puente losa apoyada L=12 m", "puente-losa-apoyada-12m", ["CAR-08.01", "CAR-08.02.01", "CAR-08.03", "CAR-08.04.01", "CAR-08.04.03", "CAR-08.04.04", "CAR-08.05", "CAR-08.06", "CAR-08.07", "CAR-05"]),
            ("Puente de vigas postensadas L=25 m", "puente-vigas-postensadas-25m", ["CAR-08"]),
        ],
    ),
    (
        "Saneamiento",
        [
            ("Redes de agua y desagüe urbanas", "redes-agua-desague", ["S-01", "S-02", "S-03", "S-04", "S-05", "S-06"]),
            ("Alcantarillado", "alcantarillado", ["S-01", "S-04", "S-05", "S-06"]),
            ("Agua potable urbana", "agua-potable", ["S-01", "S-02", "S-03", "S-05", "S-06", "S-09"]),
            ("UBS — unidad básica de saneamiento", "ubs-unidad-basica-saneamiento", ["S-07", "S-02"]),
            ("PTAR de lodos activados", "ptar-lodos-activados", ["S-01", "S-05", "S-08"]),
            ("Agua potable rural con reservorio", "agua-potable-rural-reservorio", ["S-03", "S-09", "IS-09"]),
        ],
    ),
    (
        "Hidráulica",
        [
            ("Canal de riego", "canal-riego", ["HID-"]),
            ("Canal revestido de concreto (detalle)", "canal-revestido-concreto-detalle", ["HID-"]),
            ("Bocatoma y desarenador", "bocatoma-desarenador", ["HID-"]),
            ("Defensa ribereña", "defensa-riberena", ["HID-"]),
        ],
    ),
    (
        "Habilitaciones urbanas y pavimentos",
        [
            ("Habilitación urbana", "habilitacion-urbana", ["HAB-", "P-", "S-01", "S-02", "S-03", "S-04", "IE-05.11"]),
            ("Pistas y veredas", "pistas-y-veredas", ["P-", "HAB-03"]),
            ("Pista flexible", "pista-flexible", ["P-"]),
            ("Pista rígida", "pista-rigida", ["P-"]),
            ("Veredas y sardineles", "veredas-sardineles", ["P-", "HAB-03"]),
            ("Parque urbano", "parque-urbano", ["HAB-"]),
            ("Parques, jardines y áreas verdes", "parques-jardines-y-areas-verdes", ["HAB-", "IE-05.11", "IE-05.13"]),
            ("Carretera vecinal", "carretera-vecinal", ["CAR-"]),
            ("Carretera asfaltada", "carretera-asfaltada", ["CAR-"]),
        ],
    ),
]


def set_run(run, *, size=11, bold=False, color=INK, italic=False):
    run.font.name = "Calibri"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic
    run.font.color.rgb = color


def p(doc, text, *, size=11, bold=False, color=INK, space=8, center=False, italic=False):
    para = doc.add_paragraph()
    para.paragraph_format.space_after = Pt(space)
    para.paragraph_format.space_before = Pt(0)
    if center:
        para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_run(para.add_run(text), size=size, bold=bold, color=color, italic=italic)
    return para


def bullets(doc, items, color=INK):
    for item in items:
        para = doc.add_paragraph(style="List Bullet")
        para.paragraph_format.space_after = Pt(3)
        set_run(para.add_run(item), size=11, color=color)


def heading(doc, text, level=1):
    h = doc.add_heading(text, level=level)
    for run in h.runs:
        run.font.color.rgb = NAVY if level == 1 else GOLD
        run.font.name = "Calibri"


def pausa(doc, codigo, titulo, cuerpo):
    p(doc, f"PAUSA {codigo} — {titulo}", size=12, bold=True, color=RED, space=4)
    p(doc, cuerpo, size=11, color=RED, space=10)


def tabla(doc, headers, rows, widths=None):
    t = doc.add_table(rows=1 + len(rows), cols=len(headers))
    t.style = "Table Grid"
    hdr = t.rows[0].cells
    for i, h in enumerate(headers):
        hdr[i].text = ""
        para = hdr[i].paragraphs[0]
        set_run(para.add_run(h), size=9, bold=True, color=NAVY)
    for r_i, row in enumerate(rows):
        cells = t.rows[r_i + 1].cells
        for c_i, val in enumerate(row):
            cells[c_i].text = ""
            para = cells[c_i].paragraphs[0]
            para.paragraph_format.space_after = Pt(0)
            set_run(para.add_run(str(val)), size=8, color=INK)
    if widths:
        for row in t.rows:
            for i, w in enumerate(widths):
                row.cells[i].width = Cm(w)
    p(doc, "", space=8)


def cargar_partidas():
    files = [
        LIB / "partidas.ts",
        LIB / "partidasObras.ts",
        LIB / "partidasInstalaciones.ts",
        LIB / "partidasAmpliado.ts",
        LIB / "partidasBiomedicas.ts",
        LIB / "partidasPuentesSaneamiento.ts",
        LIB / "partidasElectromecanicas.ts",
        LIB / "partidasCatalogoDetalle.ts",
        LIB / "partidasEquipamiento.ts",
    ]
    seen: dict[str, tuple] = {}
    for f in files:
        text = f.read_text(encoding="utf-8")
        for raw in P_RE.findall(text):
            esp, cap, codigo, desc, und = (unquote(x) for x in raw)
            seen[codigo] = (esp, cap, codigo, desc, und)
        for codigo, nombre, tipo, cap, upss, ambiente, _desde in EQ_RE.findall(text):
            und = "glb" if codigo.startswith("EQ-GG-") else "und"
            verbo = "" if tipo == "expediente" else "Suministro de " if tipo in ("utensilio", "herramienta") else "Suministro e instalación de "
            seen[codigo] = (
                "equipamiento",
                cap,
                codigo,
                f"{verbo}{nombre} ({tipo} · {upss} · {ambiente})",
                und,
            )
        for n, mm in THW_RE.findall(text):
            codigo = f"IE-04.10.{n}"
            seen[codigo] = (
                "electricas",
                "04 Alimentadores",
                codigo,
                f"Suministro e instalación de conductor THW-90 {mm} mm² en canalización existente",
                "m",
            )
        for n, nombre in AE_RE.findall(text):
            codigo = f"IM-09.01.{n}"
            seen[codigo] = (
                "mecanicas",
                "09 Equipamiento biomédico MINSA-DIEM",
                codigo,
                f"Suministro e instalación de {nombre} (AE MINSA-DIEM · SIGA grupo 65)",
                "und",
            )
    return [seen[k] for k in sorted(seen)]


def cargar_minsa():
    text = (LIB / "categoriasMinsa.ts").read_text(encoding="utf-8")
    ae_desde = {n: cat for n, cat in AE_DESDE_RE.findall(text)}
    cod_desde = {c: cat for c, cat in COD_DESDE_RE.findall(text)}
    eq_text = (LIB / "partidasEquipamiento.ts").read_text(encoding="utf-8")
    for codigo, _nombre, _tipo, _cap, _upss, _ambiente, desde in EQ_RE.findall(eq_text):
        cod_desde[codigo] = desde
    meta = {}
    for cat, nombre, alias, resumen, norma in META_RE.findall(text):
        meta[cat] = {"nombre": nombre, "alias": alias, "resumen": resumen, "norma": norma}
    return ae_desde, cod_desde, meta


def permitido(codigo: str, cat: str, ae_desde: dict, cod_desde: dict) -> bool:
    m = re.match(r"^IM-09\.01\.(\d{2})$", codigo)
    if m:
        desde = ae_desde.get(m.group(1))
        return bool(desde) and RANK[cat] >= RANK[desde]
    desde = cod_desde.get(codigo)
    if desde:
        return RANK[cat] >= RANK[desde]
    return True


def filtra(partidas, prefijos):
    out = []
    for row in partidas:
        codigo = row[2]
        if any(codigo.startswith(px) for px in prefijos):
            out.append(row)
    return out


def protocolo(doc):
    p(doc, "MEMORIACALC  ·  ingenieria.miacademiapreu.com  ·  Presupuestos  ·  PRE-00", size=10, bold=True, color=GOLD, center=True)
    p(doc, "Instrucción única del agente de metrados", size=22, bold=True, color=NAVY, center=True, space=4)
    p(
        doc,
        "PDF por especialidad. Primero el catálogo existente. Después la lectura. Al final, solo metrados de esos códigos. "
        "Esta es la única instrucción: protocolo + catálogo completo por especialidad + catálogo por plantilla, "
        "incluida cada categoría MINSA (I-1 a III-2) en tres alcances.",
        size=12,
        italic=True,
        color=INK,
        center=True,
        space=16,
    )

    heading(doc, "0. Cómo está armada esta instrucción")
    p(doc, "Hay una sola instrucción. No hay un «agente de hospitales» aparte ni un catálogo paralelo. El otro trabajo (categorías NTS 021, filtro de AE, plantillas minsa-i-1 … minsa-iii-2) queda integrado aquí.")
    bullets(
        doc,
        [
            "Parte I (§1–§11): protocolo. Cómo se lee el PDF y cómo se prohíbe inventar partidas.",
            "Parte II (§12): catálogo maestro. Todas las partidas del programa, por especialidad y capítulo, con código, unidad y descripción oficiales.",
            "Parte III (§13): catálogo por plantilla de proyecto (edificaciones, puentes, saneamiento, hidráulica, habilitaciones).",
            "Parte IV (§14): diez categorías MINSA × tres alcances (infraestructura + equipo, solo obra civil, solo equipamiento). Cada una con su catálogo de AE y electromecánica admitidos.",
            "Parte V (§15): matriz AE × categoría. Si la celda está vacía, ese equipo no entra en la plantilla de esa categoría.",
            "Parte VI (§16): PRE-05 especificaciones técnicas.",
        ],
    )
    p(doc, "La plantilla es semilla. El plano manda. Si el plano no muestra el ítem, no se metra aunque figure en la plantilla. Si el plano muestra un ítem y el código existe en ESTE catálogo, se metra.", bold=True)

    heading(doc, "1. Para qué sirve")
    p(
        doc,
        "El usuario entrega planos en PDF por especialidad en Presupuestos → PRE-00. El agente, o el chat de Cursor con los mismos PDF, "
        "no arma un presupuesto nuevo. Recibe el catálogo ya filtrado a esa especialidad (y, si aplica, a la categoría MINSA) y escribe el metrado de cada código que el plano justifique.",
    )

    heading(doc, "2. Cómo se entrega el expediente")
    p(doc, "Un lote = una especialidad + sus láminas. No se mezclan en el mismo lote:")
    bullets(doc, list(ESP_LABEL.values()))
    p(doc, "Si la obra es un establecimiento de salud, el usuario declara la categoría (I-1 … III-2). El catálogo de IM-08, IM-09, IE-06 e IE-07 llega ya recortado a esa categoría.")

    heading(doc, "3. Prohibiciones absolutas")
    bullets(
        doc,
        [
            "Inventar un código (ARQ-99.99.99, 01.01, NUEVA-01 o un código «parecido»).",
            "Inventar la descripción. El texto oficial es el de las tablas de esta instrucción y de PRE-01.",
            "Inventar precios, APU, rendimientos o marcas.",
            "Cambiar la unidad del catálogo para que «entre» un número.",
            "Borrar del catálogo una partida porque no se ve. El catálogo no se edita.",
            "Rellenar con una plantilla, con «lo típico» o con otro proyecto.",
            "Contar un A1 de memoria. Si no se lee, es ILEGIBLE.",
            "Convertir un hallazgo (poste de concesión, fleje, hydrobox, luminaria de marca del cliente) en línea si no tiene código en ESTE catálogo. N2XOH, splitter, ONT y perifoneo SÍ pueden tener IEM-/COM-.",
            "Usar una partida de otra especialidad aunque «encaje».",
            "Usar «no detallado» si existe un código específico (Ø, mm², W, kVA, HP, tipo de luminaria, tipo de AE).",
            "Duplicar el mismo código. Si varias láminas aportan, se suma y se explica.",
            "Redondear a decenas «por limpieza» ni añadir 5 % o 10 % de desperdicio si el plano no lo pide.",
            "Poner un tomógrafo, resonador, ventilador de UCI, planta PSA o ATS de 800 A en un I-1 / I-2 / I-3 si el plano no los dibuja y el catálogo filtrado no los trae.",
        ],
    )

    heading(doc, "4. Orden de trabajo (no se altera)")
    heading(doc, "4.1 Fase 0 — Catálogo primero", 2)
    p(doc, "Todavía no se mira ninguna imagen.")
    bullets(
        doc,
        [
            "Lee todas las filas (código | unidad | capítulo | descripción).",
            "Cuenta. El sistema declara CATÁLOGO_N.",
            "Agrupa por capítulo y memoriza la unidad legal.",
            "Marca las «no detallado»: solo si el trabajo se ve y no hay código específico.",
        ],
    )
    pausa(doc, "A1", "Comparar el catálogo", "Si el conteo ≠ CATÁLOGO_N, se detiene. Cero líneas. «PAUSA A1 FALLIDA: catálogo incompleto».")

    heading(doc, "4.2 Fase 1 — Identidad de cada lámina", 2)
    p(doc, "Del rótulo: obra, especialidad escrita, número de lámina, escala, típica/única, «igual a» / «se repite», norte, cotas, niveles, y —en salud— la categoría I-1 … III-2.")
    p(doc, "Tipo de lámina: A cuadro; B unifilar/isométrico; C leyenda; D planta; E detalle; F key plan; G notas; H ilegible.")
    pausa(doc, "B0", "Inventario de láminas", "Lista archivo, página, tipo A–H, escala y si es típica. Un recorte vacío es H, no se rellena.")

    heading(doc, "4.3 Fase 2 — Jerarquía de verdad", 2)
    bullets(
        doc,
        [
            "1. Cuadro o tabla rotulada en el plano.",
            "2. Unifilar, isométrico o esquema con calibres, diámetros o cotas.",
            "3. Leyenda (dice QUÉ, no CUÁNTOS).",
            "4. Planta o corte típico × repeticiones escritas.",
            "5. Criterio de obra, solo si 1–4 no metran, confianza baja y fórmula.",
        ],
    )
    p(doc, "Un conteo visual nunca gana a un cuadro.")

    heading(doc, "4.4 Fase 3 — Constantes de obra", 2)
    p(doc, "Solo lo escrito: torres, pisos, unidades, anchos, espesores, diámetros, mm², W, kVA, HP, f'c, y qué suministra el cliente.")
    pausa(doc, "B1", "Comparar constantes", "Sin número de repeticiones no se multiplica «por lo habitual». Se metra la unidad típica y se pide el factor en revisión.")

    heading(doc, "4.5 Fase 4 — Cómo se lee cada lámina", 2)
    bullets(
        doc,
        [
            "Cuadro: se transcribe. Un amperaje no es metros de tubería.",
            "Unifilar: calibres, polos, diámetros, tableros, pozos. Longitud solo con cota.",
            "Leyenda: diccionario. Cero cantidades.",
            "Planta: símbolos solo si no hay cuadro. Áreas con cotas, no con píxeles.",
            "Detalle: no multiplica el edificio.",
            "Key plan: solo el factor.",
            "Notas: alcance. Van a hallazgos.",
            "Ilegible: se declara.",
        ],
    )

    heading(doc, "4.6 Fase 5 — Del plano al código existente", 2)
    bullets(
        doc,
        [
            "¿Hay código cuya descripción y unidad coinciden? Se copia literal.",
            "Si hay varios (Ø20 vs Ø25, THW-90 10 vs 16 mm², LED 36 W vs 40 W, ATS 250 A vs 400 A, AE monitor vs ECG): el que el plano nombre.",
            "Si el trabajo se ve y ningún código encaja: no_catalogadas. Nunca una línea nueva.",
            "Si el código existe y la cantidad no se lee: revisión, no se lista con un número inventado.",
        ],
    )
    pausa(doc, "C1", "Comparar extraído versus asignado", "Tres listas: asignadas; vistas sin código; código existente sin cantidad. No se rellena con plantilla.")

    heading(doc, "5. Lectura por especialidad")
    p(doc, "Arquitectura: trazo, demoliciones, tierras, muros, tarrajeos, cielos, contrapisos, pisos (el formato 30×30 / 45×45 / 60×60 elige el código), zócalos, puertas por medida, ventanas, barandas, pinturas, coberturas. Pausa C-ARQ: vanos del cuadro = puertas + ventanas.")
    p(doc, "Estructuras: solado, cimiento, ciclópeo, concreto por elemento y f'c, acero fy 4200 por elemento, encofrado, ladrillo de techo. Pausa C-EST: si hay volumen y existen acero y encofrado de ese elemento, se metran ambos o se declara por qué falta uno. Sin cuantía no hay kg inventados.")
    p(doc, "Sanitarias: puntos, aparatos, PVC por Ø, cobre tipo L o K por Ø (IS-08), válvulas de bronce por Ø y tipo (IS-09, incluida la llave especial de piso). Pausa C-IS: aparatos ≈ puntos del mismo ambiente. No sustituir cobre por PVC ni al revés.")
    p(doc, "Eléctricas: cuadro de cargas > unifilar > leyenda > planta × N. Puntos, tableros, pozo, THW-90 por mm², PVC-P por Ø, luminaria por tipo y wattaje (suministro e instalación), ATS por amperaje, grupo por kVA, UPS. Si la luminaria es del cliente, no_catalogadas. Pausa C-IE: luces del cuadro × unidades + áreas comunes = puntos de iluminación.")
    p(doc, "Comunicaciones: TV, teléfono, data, intercomunicador, portería, rack. Un splitter o un poste NAP no se convierte en partida si no está en el catálogo.")
    p(doc, "Mecánicas: bombas, CI (cada ítem: Sch.40 por Ø, sprinkler por tipo, siamesa, gabinete 45/65, bombas HP, jockey, OS&Y, PIV, extintores), extractor, ascensor. Gases (IM-08) y AE oficiales (IM-09) solo si la obra es de salud y la categoría lo admite.")
    p(doc, "Equipamiento hospitalario: lote propio. Prefijos EQ- (equipo), MA- (maquinaria), UT- (utensilio), MO- (mobiliario), HE- (herramienta). Familias por UPSS: ME, DX, EM, CQ, LAB, OD, CE, CX, HO, UC, NE, FA, HD, RE; no médicos CM, LV, MT, RS, OF, AL, SS. Expediente: EQ-GG-001 a EQ-GG-008 (gastos de la dotación: movilización, transporte, instalación, capacitación, protocolos, patrimonio). No se metra con IM-09 si el plano usa el código EQ-.")
    p(doc, "Pavimentos y carreteras: espesor escrito = partida de ese espesor. Puentes: CAR-08 (ataguía, pilote Ø1.20, estribo, pila, apoyo, postensado, losa e=20 cm, junta, baranda, enrocado, prueba de carga). Pausa C-VIA: área ≈ longitud × ancho.")
    p(doc, "Saneamiento: tubería por Ø, UBS, biodigestor, PTAR (reja, desarenador, reactor, sedimentador, lecho, soplador, cloración), reservorio. Pausa C-SAN: conexiones ≤ lotes si el plano da lotes.")
    p(doc, "Hidráulica y habilitaciones: sección y espesor del plano eligen HID- y HAB-.")

    heading(doc, "6. Establecimientos de salud — NTS 021")
    p(doc, "Un I-1 no es un III-1. No tienen la misma envergadura, ni las mismas UPSS, ni las mismas especialidades, ni el mismo equipamiento. El desglose de partidas de cada uno está en el §14. Aquí solo las reglas.")
    bullets(
        doc,
        [
            "I-1 Puesto de salud: consulta externa básica. Sin internamiento, QX, UCI, lab propio, RX, TAC, RM, gases de red, ATS ni grupo.",
            "I-2 Puesto/centro con médico: médico cirujano. Laboratorio tercerizado. Sin internamiento 24 h ni quirófano.",
            "I-3 Centro de salud: consulta + patología clínica + odontología. Sin hospitalización 24 h, sin CQ, sin TAC/RM.",
            "I-4 Centro de salud con internamiento: consulta, lab, farmacia, internamiento limitado / partos. Sin UCI, sin TAC/RM, sin PSA.",
            "II-1 Hospital I: emergencia 24 h, hospitalización, CO, CQ, imágenes básicas, banco de sangre. Sin UCI completa ni TAC/RM.",
            "II-2 Hospital II: lo del II-1 + UCI, TAC, arco en C, planta de oxígeno.",
            "II-E Hospital especializado II: solo el perfil (materno, oncológico, pediátrico). No se copia un hospital general.",
            "III-1 Hospital III / nacional: alta complejidad (RM, PSA, UPS, ATS de gran amperaje).",
            "III-E Instituto especializado: alta complejidad del perfil. No se copia un III-1 general.",
            "III-2 Instituto nacional: máxima complejidad del perfil (INEN, INCOR, etc.).",
        ],
    )
    p(doc, "Tres plantillas por categoría, no una sola «hospital»:")
    bullets(
        doc,
        [
            "minsa-XX — infraestructura + equipamiento (obra civil + instalaciones + dotación EQ/MA/UT/MO/HE + gases/planta).",
            "minsa-XX-obra-civil — solo ARQ, EST, IS, IE de salidas/tableros, COM e IM de bombeo/CI/ascensor.",
            "minsa-XX-equipamiento — dotación EQ/MA/UT/MO/HE por UPSS, EQ-GG (gastos de expediente), IM-08, IE-06 e IE-07 admitidos en esa categoría. GG 12 % y utilidad 8 %.",
        ],
    )
    pausa(
        doc,
        "C-SALUD",
        "Comparar categoría versus equipos",
        "Lista la categoría, las UPSS vistas y los códigos EQ-/MA-/UT-/MO-/HE-, IM-08, IE-06 e IE-07 asignados. "
        "Si se asignó TAC, RM o PSA a un I-1 / I-2 / I-3 y el plano no los dibuja, se corrige. "
        "Lo que el plano muestre y el catálogo filtrado no tenga va a no_catalogadas.",
    )

    heading(doc, "7. Lo que el plano muestra y el catálogo no tiene")
    p(doc, "Se anota en no_catalogadas con cantidad leída. Se pide en revisión si el humano crea el código en PRE-01 y relanza. Nunca se abre una línea nueva.")

    heading(doc, "8. Criterio de obra permitido")
    bullets(
        doc,
        [
            "Tubería empotrada de iluminación o de toma: 6 a 10 m por punto, declarado.",
            "Alimentador: cota del unifilar, o 10 / 15 / 20 m si el plano lo escribe.",
            "Acero: cuantía del plano × volumen. Sin cuantía no hay kilogramos.",
        ],
    )
    p(doc, "Sin fórmula = invención = entregable inválido.")

    heading(doc, "9. Pausa D1 — Autocontrol")
    bullets(
        doc,
        [
            "Todas las líneas usan un código copiado literal de las tablas de esta instrucción / del catálogo del lote.",
            "Ningún código inventado ni «aproximado».",
            "Unidad = unidad del catálogo.",
            "Metrados > 0. Sin códigos repetidos.",
            "Las cantidades de cuadro no fueron sustituidas por un conteo visual.",
            "El factor típico está justificado o no se usó.",
            "Lo visto sin código está en no_catalogadas.",
            "No hay precios.",
            "Si hay categoría MINSA, C-SALUD está ejecutada.",
        ],
    )
    pausa(doc, "D1", "Si A1 o D1 fallan", "Cero líneas de presupuesto. Un presupuesto vacío es preferible a un presupuesto inventado.")

    heading(doc, "10. Qué recibe el humano en PRE-01")
    bullets(
        doc,
        [
            "Solo partidas del catálogo, con metrado y con el precio unitario del programa.",
            "Observaciones con rótulo, escala, especialidad, categoría MINSA y pausas.",
            "Lista de lo visto sin partida.",
            "Puntos de revisión.",
        ],
    )
    p(doc, "El agente no emite la proforma. El humano revisa, corrige y recién entonces imprime o exporta.")

    heading(doc, "11. Frase de control")
    p(
        doc,
        "«Cargué el catálogo de la especialidad (y de la categoría MINSA si aplica), conté CATÁLOGO_N partidas, leí los PDF de ese lote, "
        "asigné metrados solo a códigos existentes de las tablas de esta instrucción y dejé fuera, en no catalogadas, todo lo que el plano muestra "
        "y el programa todavía no tiene.»",
        italic=True,
    )
    p(doc, "Si el agente no puede firmar esa frase, el entregable no se usa.", bold=True, color=RED)


def parte_catalogo_maestro(doc, partidas):
    heading(doc, "12. Catálogo maestro MemoriaCalc — desglose completo")
    p(
        doc,
        f"Estas son las {len(partidas)} partidas vigentes del programa. El agente no memoriza precios. "
        "Sí debe reconocer el código, la unidad y la descripción. Si el plano pide un Ø, un mm², un wattaje, un kVA o un AE que está aquí, usa ese código. "
        "Si no está, no_catalogadas. No se resume esta lista: se usa tal cual.",
    )
    por_esp = defaultdict(lambda: defaultdict(list))
    for esp, cap, codigo, desc, und in partidas:
        por_esp[esp][cap].append((codigo, und, desc))

    n_sec = 0
    for esp, label in ESP_LABEL.items():
        caps = por_esp.get(esp)
        if not caps:
            continue
        n_sec += 1
        total = sum(len(v) for v in caps.values())
        heading(doc, f"12.{n_sec} {label} — {total} partidas", 2)
        for cap, filas in caps.items():
            heading(doc, f"{cap} ({len(filas)})", 3)
            tabla(doc, ["Código", "Und", "Descripción oficial"], [(a, b, c) for a, b, c in filas], widths=[3.4, 1.6, 12.0])


def parte_plantillas(doc, partidas):
    heading(doc, "13. Catálogos de partidas por plantilla (proyecto)")
    p(
        doc,
        "Cada plantilla de PRE-01 tiene un catálogo cerrado: las partidas que ese tipo de obra PUEDE usar. "
        "No es un metrado. El agente no copia las cantidades de la plantilla. Lee el PDF y metra solo los códigos de este subconjunto que el plano justifique.",
    )
    idx = 0
    for grupo, items in PLANTILLA_CATALOGO:
        idx += 1
        heading(doc, f"13.{idx} {grupo}", 2)
        for nombre, _pid, prefs in items:
            filas = filtra(partidas, prefs)
            heading(doc, f"{nombre} — {len(filas)} partidas posibles", 3)
            p(doc, f"Prefijos de catálogo: {', '.join(prefs)}. Si el plano muestra un trabajo fuera de estos prefijos, no_catalogadas (o se procesa en el lote de otra especialidad).")
            tabla(doc, ["Código", "Und", "Descripción oficial"], [(a[2], a[4], a[3]) for a in filas], widths=[3.4, 1.6, 12.0])


def parte_minsa(doc, partidas, ae_desde, cod_desde, meta):
    heading(doc, "14. Catálogos MINSA por categoría y por alcance")
    p(
        doc,
        "NTS N° 021-MINSA/DGSP-V.03 (RM 546-2011), NTS 113 (1.er nivel), NTS 110 (2.º nivel) y NTS 119 (3.er nivel). "
        "Coordinado con el módulo de categorías del programa: el catálogo de equipos se filtra con la misma regla (categoría mínima del código ≤ categoría de la obra).",
    )
    p(
        doc,
        "La obra civil (ARQ, EST, IS de redes y aparatos, IE de salidas y tableros, COM, IM de bombeo/CI/ascensor) usa el catálogo maestro de esas familias. "
        "Lo que cambia entre I-1 y III-2 es el EQUIPAMIENTO y la ELECTROMECÁNICA DE PLANTA. Eso es lo que se detalla abajo, tres veces por categoría.",
    )

    eq_codes = [r for r in partidas if re.match(r"^(IM-08\.|IM-09\.|IE-06\.|IE-07\.|EQ-|MA-|UT-|MO-|HE-)", r[2])]

    for i, cat in enumerate(CATS, start=1):
        m = meta.get(cat, {"nombre": cat, "alias": cat, "resumen": "", "norma": ""})
        admitidos = [r for r in eq_codes if permitido(r[2], cat, ae_desde, cod_desde)]
        ae = [r for r in admitidos if r[2].startswith("IM-09.")]
        dot = [r for r in admitidos if re.match(r"^(EQ-|MA-|UT-|MO-|HE-)", r[2])]
        gases = [r for r in admitidos if r[2].startswith("IM-08.")]
        ats = [r for r in admitidos if r[2].startswith("IE-06.")]
        ge = [r for r in admitidos if r[2].startswith("IE-07.")]
        heading(doc, f"14.{i} {m.get('nombre', cat)}", 2)
        p(doc, m.get("resumen", ""), italic=True)
        p(doc, m.get("norma", ""), size=10, color=GOLD)
        p(doc, f"Plantillas en PRE-01: minsa-{cat.lower()} (infra + equipo) · minsa-{cat.lower()}-obra-civil · minsa-{cat.lower()}-equipamiento.")

        heading(doc, f"{cat} — alcance 1: infraestructura + equipamiento", 3)
        p(doc, "Usa: (a) catálogo maestro de ARQ, EST, IS, IE de salidas/alimentadores/luminarias, COM e IM de bombeo/CI/ascensor, según lo que el plano dibuje; (b) el catálogo de equipo de esta categoría, que sigue.")
        heading(doc, f"{cat} — alcance 2: solo obra civil e instalaciones", 3)
        p(doc, "Prohibido metrar EQ-/MA-/UT-/MO-/HE-, IM-08, IM-09, IE-06 e IE-07. Si el plano de arquitectura o de IE muestra un grupo o un ATS, va a no_catalogadas de ESE lote y se procesa en el lote de equipamiento o se pide crear el código.")
        heading(doc, f"{cat} — alcance 3: solo equipamiento", 3)
        p(doc, f"Familias de este alcance: EQ/MA/UT/MO/HE por UPSS, EQ-GG, IM-08, IE-06 e IE-07 ({len(admitidos)} códigos). Incluye gastos generales de la dotación (EQ-GG) y, en el presupuesto, GG 12 % + utilidad 8 % + IGV.")

        if dot:
            heading(doc, f"Dotación por UPSS admitida en {cat} ({len(dot)})", 3)
            tabla(doc, ["Código", "Und", "Equipo / elemento (tipo · UPSS · ambiente)"], [(a[2], a[4], a[3]) for a in dot], widths=[3.4, 1.6, 12.0])
        else:
            p(doc, "Esta categoría no admite dotación EQ en el catálogo filtrado.", color=RED)

        if ae:
            heading(doc, f"AE MINSA-DIEM admitidos en {cat} ({len(ae)})", 3)
            tabla(doc, ["Código", "Und", "Equipo (suministro e instalación)"], [(a[2], a[4], a[3]) for a in ae], widths=[3.4, 1.6, 12.0])
        else:
            p(doc, "Esta categoría no admite AE de planta en el catálogo filtrado.", color=RED)

        if gases:
            heading(doc, f"Gases medicinales y planta de O2 admitidos en {cat} ({len(gases)})", 3)
            tabla(doc, ["Código", "Und", "Descripción oficial"], [(a[2], a[4], a[3]) for a in gases], widths=[3.4, 1.6, 12.0])
        else:
            p(doc, f"{cat} no admite red de gases ni planta PSA en plantilla. Si el plano las dibuja, no_catalogadas.", color=RED)

        if ats:
            heading(doc, f"ATS / UPS admitidos en {cat} ({len(ats)})", 3)
            tabla(doc, ["Código", "Und", "Descripción oficial"], [(a[2], a[4], a[3]) for a in ats], widths=[3.4, 1.6, 12.0])
        else:
            p(doc, f"{cat} no admite ATS ni UPS en plantilla.", color=RED)

        if ge:
            heading(doc, f"Grupo electrógeno admitido en {cat} ({len(ge)})", 3)
            tabla(doc, ["Código", "Und", "Descripción oficial"], [(a[2], a[4], a[3]) for a in ge], widths=[3.4, 1.6, 12.0])
        else:
            p(doc, f"{cat} no admite grupo electrógeno en plantilla.", color=RED)


def parte_matriz(doc, partidas, ae_desde):
    heading(doc, "15. Matriz AE MINSA × categoría")
    p(doc, "Marca = el código entra en el catálogo filtrado de esa categoría (y de las superiores). Vacío = no entra. Si el plano de un I-1 muestra un tomógrafo, no se usa IM-09.01.27: va a no_catalogadas.")
    aes = [r for r in partidas if r[2].startswith("IM-09.01.")]
    headers = ["Código", "Equipo"] + CATS
    rows = []
    for _esp, _cap, codigo, desc, _und in aes:
        n = codigo.split(".")[-1]
        nombre = desc.replace("Suministro e instalación de ", "").replace(" (AE MINSA-DIEM · SIGA grupo 65)", "")
        marca = []
        for cat in CATS:
            desde = ae_desde.get(n)
            marca.append("●" if desde and RANK[cat] >= RANK[desde] else "")
        rows.append((codigo, nombre, *marca))
    tabla(doc, headers, rows, widths=[2.6, 5.2] + [0.9] * 10)


def parte_et(doc):
    heading(doc, "16. PRE-05 Especificaciones técnicas")
    p(doc, "Debajo de Cronograma de obra está PRE-05 Especificaciones técnicas. Cada partida de las tablas anteriores tiene ficha de expediente: descripción, tabla APU real (mano de obra, materiales, maquinaria y equipo del catálogo), procedimiento, medición, bases de pago, control de aceptación y normas (RNE, CNE, IEC, TIA-568, NFPA, MINSA-DIEM, EG-2013, según familia).")
    p(doc, "Las fichas desglosan el presupuesto: no inventan insumos. El material y el equipo salen de la receta APU de esa partida. Se descargan / imprimen con el presupuesto. El agente de metrados no redacta la ET ni crea partidas.")
    p(doc, "La especialidad electromecánica (IEM-) es lote propio: subestación, MCC, ATS, grupo, UPS, motores, izaje, tierra de planta, NYY/N2XSY/N2XOH, bandejas y fotovoltaica. No se metra con códigos IE- ni IM- en un lote IEM.")
    p(doc, "MemoriaCalc · PRE-00 · Instrucción única · no crear partidas · catálogo por especialidad y por plantilla · MINSA I-1 a III-2", size=9, color=GOLD, center=True, space=18)


def build():
    partidas = cargar_partidas()
    ae_desde, cod_desde, meta = cargar_minsa()
    doc = Document()
    for section in doc.sections:
        section.top_margin = Cm(1.6)
        section.bottom_margin = Cm(1.6)
        section.left_margin = Cm(1.8)
        section.right_margin = Cm(1.8)

    protocolo(doc)
    parte_catalogo_maestro(doc, partidas)
    parte_plantillas(doc, partidas)
    parte_minsa(doc, partidas, ae_desde, cod_desde, meta)
    parte_matriz(doc, partidas, ae_desde)
    parte_et(doc)

    for dest in OUTS:
        dest.parent.mkdir(parents=True, exist_ok=True)
        doc.save(dest)
        print("OK", dest, dest.stat().st_size, "partidas", len(partidas))


if __name__ == "__main__":
    build()
