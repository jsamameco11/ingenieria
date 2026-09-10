import json
import re
import sys

import openpyxl

sys.stdout.reconfigure(encoding="utf-8")

src = r"d:\Renzo\RENZO\6. EXCELS\1. ESTRUCTURAS\1. EDIFIC\4. EDIFICACIONES COMPLETO\ANA DISEÑO EDIFI.xlsx"
wb = openpyxl.load_workbook(src, data_only=True)
ws = wb["Cuadros E030"]
out = r"C:\Users\Renzo\Desktop\WEB MEMORIAS DESCRIPTIVAS\scripts\e030-inspect.txt"

lines = []


def p(*a):
    lines.append(" ".join(str(x) if x is not None else "" for x in a))


p("max_row", ws.max_row, "max_col", ws.max_column)
p("=== defined names ===")
names = []
for name, dn in wb.defined_names.items():
    attr = str(dn.attr_text or "")
    if "Cuadros E030" in attr:
        names.append((name, attr))
p("count", len(names))
for name, attr in names[:80]:
    p(name, "=>", attr)

p("\n=== A-E rows 1-80 ===")
for r in range(1, 81):
    vals = [ws.cell(r, c).value for c in range(1, 6)]
    if any(v is not None for v in vals):
        p(r, vals)

p("\n=== A-E rows 1800-1900 ===")
for r in range(1800, 1901):
    vals = [ws.cell(r, c).value for c in range(1, 6)]
    if any(v is not None for v in vals):
        p(r, vals)

p("\n=== unique D values among C-looking districts ===")
from collections import Counter

c = Counter()
for r in range(1, ws.max_row + 1):
    v = ws.cell(r, 3).value
    d = ws.cell(r, 4).value
    if isinstance(v, str) and re.match(r"^\d+\s+", v.strip()):
        c[repr(d)] += 1
p(c.most_common(20))
p("district-like rows", sum(c.values()))

p("\n=== sample district rows ===")
n = 0
for r in range(1, ws.max_row + 1):
    v = ws.cell(r, 3).value
    if isinstance(v, str) and re.match(r"^\d+\s+", v.strip()):
        p(r, ws.cell(r, 1).value, ws.cell(r, 2).value, v, ws.cell(r, 4).value, ws.cell(r, 5).value)
        n += 1
        if n >= 25:
            break

p("\n=== ANALISIS sheet location cells ===")
if "ANALISIS SISMICO ESTATICO X-X" in wb.sheetnames:
    ax = wb["ANALISIS SISMICO ESTATICO X-X"]
    for r in range(1, 40):
        vals = [ax.cell(r, c).value for c in range(1, 10)]
        if any(v is not None for v in vals):
            p(r, vals)

with open(out, "w", encoding="utf-8") as f:
    f.write("\n".join(lines))
print("wrote", out, "lines", len(lines))
