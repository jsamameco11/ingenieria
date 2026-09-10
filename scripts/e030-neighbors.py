import sys
import openpyxl

sys.stdout.reconfigure(encoding="utf-8")
src = r"d:\Renzo\RENZO\6. EXCELS\1. ESTRUCTURAS\1. EDIFIC\4. EDIFICACIONES COMPLETO\ANA DISEÑO EDIFI.xlsx"
wb = openpyxl.load_workbook(src, data_only=True)
ws = wb["Cuadros E030"]
for start, end in [(300, 320), (870, 900), (1080, 1100), (1610, 1630)]:
    print(f"\n===== {start}-{end} =====")
    for r in range(start, end + 1):
        print(
            r,
            ws.cell(r, 1).value,
            "|",
            ws.cell(r, 2).value,
            "|",
            ws.cell(r, 3).value,
            "|",
            ws.cell(r, 4).value,
            "|",
            ws.cell(r, 5).value,
        )
