# -*- coding: utf-8 -*-
import json
from pathlib import Path
import openpyxl

path = Path(r"D:\Renzo\RENZO\6. EXCELS\12. HIDRAULICA\2. DRENAJE PLUVIAL\DISEÑO CANALETA.xlsx")
wb = openpyxl.load_workbook(path, data_only=True)
ws = wb["C - 1 R"]
print("max_row", ws.max_row, "max_col", ws.max_column)
lines = []
for r in range(180, min(ws.max_row, 280) + 1):
    vals = []
    for c in range(1, min(ws.max_column, 16) + 1):
        v = ws.cell(r, c).value
        if v is None:
            vals.append("")
        elif isinstance(v, float):
            vals.append(f"{v:.8g}")
        else:
            vals.append(str(v).strip())
    if any(vals):
        lines.append(f"{r:3d} | " + " | ".join(vals[:14]))
text = "\n".join(lines)
Path(r"C:\Users\Renzo\Desktop\WEB MEMORIAS DESCRIPTIVAS\_canaleta_tail.txt").write_text(text, encoding="utf-8")
print(text[:12000])

# extract images
wb2 = openpyxl.load_workbook(path)
img_dir = Path(r"C:\Users\Renzo\Desktop\WEB MEMORIAS DESCRIPTIVAS\_imgs")
img_dir.mkdir(exist_ok=True)
n = 0
for name in wb2.sheetnames:
    ws = wb2[name]
    imgs = getattr(ws, "_images", []) or []
    for i, img in enumerate(imgs):
        try:
            data = img._data()
            ext = getattr(img, "format", None) or "png"
            if isinstance(ext, str) and ext.lower() in ("jpeg", "jpg"):
                ext = "jpg"
            else:
                ext = "png"
            out = img_dir / f"{name.replace(' ','_')}_{i}.{ext}"
            out.write_bytes(data)
            n += 1
            print("saved", out.name, len(data))
        except Exception as e:
            print("img fail", name, i, type(e).__name__, e)
print("images saved", n)
